import os
import glob
import json
import sys
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path
from scipy import stats
from statsmodels.stats.contingency_tables import mcnemar


TIERS = ["T1", "T2", "T3", "T4"]

MATCHED_PAIRS = [
    ("mistral-7b-instruct", "biomistral-7b",  "7B"),
    ("llama-3-8b-instruct", "openbiollm-8b",  "8B"),
    ("gemma-3-27b",         "medgemma-27b",   "27B"),
]


def acc_ci(correct_series):
    n = len(correct_series)
    k = correct_series.sum()
    if n == 0:
        return np.nan, np.nan, np.nan
    p      = k / n
    z      = 1.96
    denom  = 1 + z**2 / n
    centre = (p + z**2 / (2 * n)) / denom
    half   = (z * np.sqrt(p * (1 - p) / n + z**2 / (4 * n**2))) / denom
    return round(p * 100, 2), round((centre - half) * 100, 2), round((centre + half) * 100, 2)


def fmt_ci(acc, lo, hi):
    if np.isnan(acc):
        return "N/A"
    return f"{acc:.1f}% ({lo:.1f}–{hi:.1f})"


def load_results(results_folder):
    all_rows = []
    for f in sorted(glob.glob(f"{results_folder}/results_*.jsonl")):
        with open(f) as fh:
            for line in fh:
                try:
                    all_rows.append(json.loads(line.strip()))
                except Exception:
                    pass

    df = pd.DataFrame(all_rows)
    df["is_correct"] = df["is_correct"].astype(bool)
    if "new_tier" in df.columns:
        df["tier"] = df["new_tier"].fillna(df["tier"])
    print(f"Loaded {len(df):,} rows, {df['model'].nunique()} models")
    return df


def table1_main_results(df, output_folder):
    rows = []
    for model in sorted(df["model"].unique()):
        mdf = df[df["model"] == model]
        cat = mdf["category"].iloc[0] if "category" in mdf.columns else "unknown"
        a, lo, hi = acc_ci(mdf["is_correct"])
        row = {
            "Model":    model,
            "Category": cat,
            "N_total":  len(mdf),
            "Overall":  a,
            "Overall_CI_lower": lo,
            "Overall_CI_upper": hi,
            "Overall_fmt": fmt_ci(a, lo, hi),
        }
        for tier in TIERS:
            tdf = mdf[mdf["tier"] == tier]
            ta, tlo, thi = acc_ci(tdf["is_correct"])
            row[tier]               = ta
            row[f"{tier}_CI_lower"] = tlo
            row[f"{tier}_CI_upper"] = thi
            row[f"{tier}_N"]        = len(tdf)
            row[f"{tier}_fmt"]      = fmt_ci(ta, tlo, thi)
        rows.append(row)

    out = pd.DataFrame(rows).sort_values("Overall", ascending=False).reset_index(drop=True)
    out.to_csv(f"{output_folder}/table1_main_results.csv", index=False)
    print("Saved table1_main_results.csv")

    cols = ["Model", "Category", "Overall_fmt", "T1_fmt", "T2_fmt", "T3_fmt", "T4_fmt"]
    print("\nTABLE 1 — MAIN RESULTS")
    print("=" * 90)
    print(out[cols].to_string(index=False))
    return out


def table2_crossing_degradation(df, output_folder):
    max_cross = int(df["crossing_count"].max()) if "crossing_count" in df.columns else 9
    rows = []
    for model in sorted(df["model"].unique()):
        mdf = df[df["model"] == model]
        row = {"Model": model}
        for c in range(2, max_cross + 1):
            cdf = mdf[mdf["crossing_count"] == c]
            if len(cdf) >= 5:
                a, lo, hi = acc_ci(cdf["is_correct"])
                row[f"cross_{c}"]     = a
                row[f"cross_{c}_N"]   = len(cdf)
                row[f"cross_{c}_fmt"] = fmt_ci(a, lo, hi)
            else:
                row[f"cross_{c}"]     = np.nan
                row[f"cross_{c}_N"]   = len(cdf)
                row[f"cross_{c}_fmt"] = f"N/A (n={len(cdf)})"
        rows.append(row)

    out = pd.DataFrame(rows)
    out.to_csv(f"{output_folder}/table2_crossing_degradation.csv", index=False)
    print("Saved table2_crossing_degradation.csv")
    return out, max_cross


def table3_finetuning_paradox(df, output_folder):
    rows = []
    for general_model, bio_model, scale in MATCHED_PAIRS:
        gen_df = df[df["model"] == general_model].set_index("item_id")
        bio_df = df[df["model"] == bio_model].set_index("item_id")

        common_ids  = gen_df.index.intersection(bio_df.index)
        gen_correct = gen_df.loc[common_ids, "is_correct"]
        bio_correct = bio_df.loc[common_ids, "is_correct"]

        both_correct = (gen_correct & bio_correct).sum()
        only_gen     = (gen_correct & ~bio_correct).sum()
        only_bio     = (~gen_correct & bio_correct).sum()
        both_wrong   = (~gen_correct & ~bio_correct).sum()

        table = [[int(both_correct), int(only_gen)],
                 [int(only_bio),     int(both_wrong)]]

        result  = mcnemar(table, exact=False, correction=True)
        p_value = result.pvalue
        p_fmt   = f"{p_value:.4f}" if p_value >= 0.0001 else "<0.0001"
        sig     = ("***" if p_value < 0.001 else
                   "**"  if p_value < 0.01  else
                   "*"   if p_value < 0.05  else "ns")

        gen_acc, gen_lo, gen_hi = acc_ci(gen_correct)
        bio_acc, bio_lo, bio_hi = acc_ci(bio_correct)
        gap = round(bio_acc - gen_acc, 2) if not np.isnan(gen_acc) else np.nan

        tier_data = {}
        for tier in TIERS:
            merged = pd.DataFrame({
                "gen":  gen_correct,
                "bio":  bio_correct,
                "tier": gen_df.loc[common_ids, "tier"],
            })
            t = merged[merged["tier"] == tier]
            ga, _, _ = acc_ci(t["gen"]) if len(t) > 0 else (np.nan, None, None)
            ba, _, _ = acc_ci(t["bio"]) if len(t) > 0 else (np.nan, None, None)
            tier_data[f"{tier}_general"]    = ga
            tier_data[f"{tier}_biomedical"] = ba
            tier_data[f"{tier}_gap"]        = round(ba - ga, 2) if not np.isnan(ga) else np.nan

        rows.append({
            "Scale":            scale,
            "General_Model":    general_model,
            "Biomedical_Model": bio_model,
            "N_items":          len(common_ids),
            "General_Acc":      gen_acc,
            "General_CI":       fmt_ci(gen_acc, gen_lo, gen_hi),
            "Biomedical_Acc":   bio_acc,
            "Biomedical_CI":    fmt_ci(bio_acc, bio_lo, bio_hi),
            "Gap":              gap,
            "McNemar_p":        p_fmt,
            "Significance":     sig,
            "Both_correct":     int(both_correct),
            "Only_general":     int(only_gen),
            "Only_biomedical":  int(only_bio),
            "Both_wrong":       int(both_wrong),
            **tier_data,
        })
        print(f"  {scale}: {general_model} {gen_acc:.1f}% vs {bio_model} {bio_acc:.1f}% | gap={gap:+.1f}% | p={p_fmt} {sig}")

    out = pd.DataFrame(rows)
    out.to_csv(f"{output_folder}/table3_finetuning_paradox.csv", index=False)
    print("Saved table3_finetuning_paradox.csv")
    return out


def table4_paradox_by_tier(paradox_df, output_folder):
    rows = []
    for _, prow in paradox_df.iterrows():
        for tier in TIERS:
            rows.append({
                "Scale":            prow["Scale"],
                "Tier":             tier,
                "General_Model":    prow["General_Model"],
                "Biomedical_Model": prow["Biomedical_Model"],
                "General_Acc":      prow.get(f"{tier}_general", np.nan),
                "Biomedical_Acc":   prow.get(f"{tier}_biomedical", np.nan),
                "Gap":              prow.get(f"{tier}_gap", np.nan),
            })

    out = pd.DataFrame(rows)
    out.to_csv(f"{output_folder}/table4_paradox_by_tier.csv", index=False)
    print("Saved table4_paradox_by_tier.csv")
    return out


def table5_confidence_intervals(df, output_folder):
    rows = []
    for model in sorted(df["model"].unique()):
        mdf = df[df["model"] == model]
        cat = mdf["category"].iloc[0] if "category" in mdf.columns else "unknown"
        a, lo, hi = acc_ci(mdf["is_correct"])
        row = {
            "Model": model, "Category": cat,
            "Overall_acc": a, "CI_lower": lo, "CI_upper": hi,
            "CI_width": round(hi - lo, 2), "N": len(mdf),
        }
        for tier in TIERS:
            tdf = mdf[mdf["tier"] == tier]
            ta, tlo, thi = acc_ci(tdf["is_correct"])
            row[f"{tier}_acc"]   = ta
            row[f"{tier}_lower"] = tlo
            row[f"{tier}_upper"] = thi
            row[f"{tier}_N"]     = len(tdf)
        rows.append(row)

    out = pd.DataFrame(rows).sort_values("Overall_acc", ascending=False).reset_index(drop=True)
    out.to_csv(f"{output_folder}/table5_confidence_intervals.csv", index=False)
    print("Saved table5_confidence_intervals.csv")
    return out


def build_excel(summary_df, crossing_df, paradox_df, tier_paradox_df,
                ci_df, max_cross, output_folder):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    NAVY   = PatternFill("solid", start_color="1E2A44")
    TEAL   = PatternFill("solid", start_color="008080")
    ALT    = PatternFill("solid", start_color="EEF1F4")
    FRO    = PatternFill("solid", start_color="E8F4F8")
    BIO    = PatternFill("solid", start_color="FFF3E0")
    GEN    = PatternFill("solid", start_color="F1F8E9")
    WHITE  = PatternFill("solid", start_color="FFFFFF")
    SIG    = PatternFill("solid", start_color="E8F5E9")

    HDR_F  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    TTL_F  = Font(name="Arial", bold=True, color="1E2A44", size=13)
    SUB_F  = Font(name="Arial", italic=True, color="555555", size=10)
    BODY_F = Font(name="Arial", size=10)
    BOLD_F = Font(name="Arial", bold=True, size=10)

    thin   = Side(style="thin", color="CCCCCC")
    BDR    = Border(left=thin, right=thin, top=thin, bottom=thin)
    CTR    = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LFT    = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    def hdr(cell, fill=None):
        cell.font = HDR_F; cell.fill = fill or NAVY
        cell.alignment = CTR; cell.border = BDR

    def body(cell, fill=None, bold=False, left=False):
        cell.font = BOLD_F if bold else BODY_F
        cell.fill = fill or WHITE
        cell.alignment = LFT if left else CTR
        cell.border = BDR

    def cat_fill(cat):
        if "frontier" in str(cat): return FRO
        if "biomedical" in str(cat): return BIO
        if "general" in str(cat): return GEN
        return WHITE

    def title(ws, t, sub=""):
        ws["A1"] = t; ws["A1"].font = TTL_F; ws["A1"].alignment = LFT
        if sub:
            ws["A2"] = sub; ws["A2"].font = SUB_F; ws["A2"].alignment = LFT

    # Sheet 1
    ws1 = wb.create_sheet("Table1_Main_Results")
    title(ws1, "Table 1 — Accuracy by Model × Tier",
          "95% Wilson confidence intervals. Sorted by overall accuracy.")
    hdrs = ["Model", "Category", "N", "Overall (95% CI)",
            "T1 (95% CI)", "T2 (95% CI)", "T3 (95% CI)", "T4 (95% CI)"]
    for c, h in enumerate(hdrs, 1):
        hdr(ws1.cell(row=4, column=c, value=h))
    for ri, (_, row) in enumerate(summary_df.iterrows(), 5):
        fill = cat_fill(row["Category"])
        for ci, val in enumerate([row["Model"], row["Category"], row["N_total"],
                                   row["Overall_fmt"], row["T1_fmt"], row["T2_fmt"],
                                   row["T3_fmt"], row["T4_fmt"]], 1):
            body(ws1.cell(row=ri, column=ci, value=val),
                 fill=fill, bold=(ci == 1), left=(ci <= 2))
    for col, w in zip(["A","B","C","D","E","F","G","H"], [28,20,8,22,22,22,22,22]):
        ws1.column_dimensions[col].width = w
    ws1.freeze_panes = "A5"

    # Sheet 2
    cross_cols = [c for c in range(2, max_cross + 1)
                  if f"cross_{c}_fmt" in crossing_df.columns]
    ws2 = wb.create_sheet("Table2_Crossing_Degradation")
    title(ws2, "Table 2 — Accuracy by Cross-Domain Crossing Count")
    for c, h in enumerate(["Model"] + [f"Crossings={c}" for c in cross_cols], 1):
        hdr(ws2.cell(row=4, column=c, value=h), fill=TEAL if c > 1 else NAVY)
    for ri, (_, row) in enumerate(crossing_df.iterrows(), 5):
        fill = ALT if ri % 2 == 0 else WHITE
        body(ws2.cell(row=ri, column=1, value=row["Model"]),
             fill=fill, bold=True, left=True)
        for ci, c in enumerate(cross_cols, 2):
            body(ws2.cell(row=ri, column=ci, value=row.get(f"cross_{c}_fmt", "N/A")),
                 fill=fill)
    ws2.column_dimensions["A"].width = 28
    for i in range(2, len(cross_cols) + 2):
        ws2.column_dimensions[get_column_letter(i)].width = 18
    ws2.freeze_panes = "B5"

    # Sheet 3
    ws3 = wb.create_sheet("Table3_FinetuningParadox")
    title(ws3, "Table 3 — Fine-Tuning Paradox: Matched Pair Comparisons",
          "*** p<0.001  ** p<0.01  * p<0.05  ns = not significant")
    hdrs3 = ["Scale", "General Model", "General Acc (95% CI)",
             "Biomedical Model", "Biomedical Acc (95% CI)", "Gap (pp)", "McNemar p", "Sig."]
    for c, h in enumerate(hdrs3, 1):
        hdr(ws3.cell(row=4, column=c, value=h))
    for ri, (_, row) in enumerate(paradox_df.iterrows(), 5):
        gap  = row["Gap"]
        fill = SIG if (not np.isnan(gap) and gap < 0) else WHITE
        for ci, val in enumerate([row["Scale"], row["General_Model"], row["General_CI"],
                                   row["Biomedical_Model"], row["Biomedical_CI"],
                                   f"{gap:+.1f}" if not np.isnan(gap) else "N/A",
                                   row["McNemar_p"], row["Significance"]], 1):
            body(ws3.cell(row=ri, column=ci, value=val),
                 fill=fill, bold=(ci == 1), left=(ci in [1, 2, 4]))
    for col, w in zip(["A","B","C","D","E","F","G","H"], [8,25,22,25,22,10,12,8]):
        ws3.column_dimensions[col].width = w
    ws3.freeze_panes = "A5"

    # Sheet 4
    ws4 = wb.create_sheet("Table4_Paradox_by_Tier")
    title(ws4, "Table 4 — Fine-Tuning Paradox by Difficulty Tier")
    hdrs4 = ["Scale", "Tier", "General Model", "General Acc %",
             "Biomedical Model", "Biomedical Acc %", "Gap (pp)"]
    for c, h in enumerate(hdrs4, 1):
        hdr(ws4.cell(row=4, column=c, value=h))
    for ri, (_, row) in enumerate(tier_paradox_df.iterrows(), 5):
        gap  = row["Gap"]
        fill = SIG if (not np.isnan(gap) and gap < 0) else (ALT if ri % 2 == 0 else WHITE)
        for ci, val in enumerate([
            row["Scale"], row["Tier"], row["General_Model"],
            f"{row['General_Acc']:.1f}%" if not np.isnan(row["General_Acc"]) else "N/A",
            row["Biomedical_Model"],
            f"{row['Biomedical_Acc']:.1f}%" if not np.isnan(row["Biomedical_Acc"]) else "N/A",
            f"{gap:+.1f}" if not np.isnan(gap) else "N/A",
        ], 1):
            body(ws4.cell(row=ri, column=ci, value=val),
                 fill=fill, bold=(ci == 1), left=(ci in [1, 3, 5]))
    for col, w in zip(["A","B","C","D","E","F","G"], [8,6,25,15,25,18,10]):
        ws4.column_dimensions[col].width = w
    ws4.freeze_panes = "A5"

    # Sheet 5
    ws5 = wb.create_sheet("Table5_Confidence_Intervals")
    title(ws5, "Table 5 — Full Confidence Intervals (95% Wilson)")
    ci_hdrs = ["Model", "Category", "N", "Overall%", "CI_lower", "CI_upper", "CI_width"]
    for tier in TIERS:
        ci_hdrs += [f"{tier}%", f"{tier}_lower", f"{tier}_upper", f"{tier}_N"]
    for c, h in enumerate(ci_hdrs, 1):
        hdr(ws5.cell(row=4, column=c, value=h))
    for ri, (_, row) in enumerate(ci_df.iterrows(), 5):
        fill = cat_fill(row["Category"])
        vals = [row["Model"], row["Category"], row["N"],
                row["Overall_acc"], row["CI_lower"], row["CI_upper"], row["CI_width"]]
        for tier in TIERS:
            vals += [row[f"{tier}_acc"], row[f"{tier}_lower"],
                     row[f"{tier}_upper"], row[f"{tier}_N"]]
        for ci, val in enumerate(vals, 1):
            body(ws5.cell(row=ri, column=ci,
                          value=val if not (isinstance(val, float) and np.isnan(val)) else ""),
                 fill=fill, bold=(ci == 1), left=(ci <= 2))
    ws5.column_dimensions["A"].width = 28
    ws5.column_dimensions["B"].width = 20
    for i in range(3, len(ci_hdrs) + 1):
        ws5.column_dimensions[get_column_letter(i)].width = 12
    ws5.freeze_panes = "A5"

    path = f"{output_folder}/BioChem_XDR_All_Results.xlsx"
    wb.save(path)
    print(f"Saved BioChem_XDR_All_Results.xlsx")


if __name__ == "__main__":
    results_folder = sys.argv[1]
    output_folder  = sys.argv[2] if len(sys.argv) > 2 else "analysis_outputs"

    Path(output_folder).mkdir(parents=True, exist_ok=True)

    df = load_results(results_folder)

    summary_df     = table1_main_results(df, output_folder)
    crossing_df, max_cross = table2_crossing_degradation(df, output_folder)
    paradox_df     = table3_finetuning_paradox(df, output_folder)
    tier_paradox_df = table4_paradox_by_tier(paradox_df, output_folder)
    ci_df          = table5_confidence_intervals(df, output_folder)

    build_excel(summary_df, crossing_df, paradox_df, tier_paradox_df,
                ci_df, max_cross, output_folder)

    print("\nAnalysis complete.")
    print(f"Output: {output_folder}")